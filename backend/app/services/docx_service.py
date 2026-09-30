import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.schemas.document import ReconstructedDocument


def render_docx(reconstructed: ReconstructedDocument) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 12)):
        styles[name].font.name = "Times New Roman"
        styles[name].font.size = Pt(size)

    document.add_heading(reconstructed.title, 0)
    for section_data in reconstructed.sections:
        if section_data.main_title:
            document.add_heading(section_data.main_title, 1)

        numbered_counter = 1
        for element in section_data.elements:
            text = element.text or ""
            if element.type == "subtitle":
                document.add_heading(text, 2)
                for child in element.children:
                    document.add_paragraph(child.text or "", style="List Bullet 2" if child.type == "sub_bullet" else "List Bullet")
            elif element.type == "bullet":
                document.add_paragraph(text, style="List Bullet")
            elif element.type == "sub_bullet":
                document.add_paragraph(text, style="List Bullet 2")
            elif element.type == "numbered":
                number = element.data.get("num", numbered_counter)
                document.add_paragraph(f"{number}. {text}")
                numbered_counter = int(number) + 1 if isinstance(number, int) else numbered_counter + 1
            elif element.type == "table":
                headers = element.data.get("headers", [])
                rows = element.data.get("rows", [])
                if text:
                    document.add_heading(text, 2)
                width = max([len(headers), *(len(row) for row in rows if isinstance(row, list))], default=0)
                if width:
                    table = document.add_table(rows=1 if headers else 0, cols=width)
                    table.style = "Table Grid"
                    table.autofit = False
                    if headers:
                        for index, value in enumerate(headers[:width]):
                            table.rows[0].cells[index].text = str(value)
                    for values in rows:
                        if isinstance(values, list):
                            cells = table.add_row().cells
                            for index, value in enumerate(values[:width]):
                                cells[index].text = str(value)
            elif element.type == "formula":
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run(text).italic = True
            elif text:
                document.add_paragraph(text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
